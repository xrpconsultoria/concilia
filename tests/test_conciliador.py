from datetime import date
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
import sqlite3
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from pypdf import PdfWriter

from app import create_app
from conciliador.domain import InputError, Movement, Statement, classify, money
from conciliador.makro import HEADERS, ROOT, export_workbook, learn_workbook, prepare_private_chart, validate_code
from conciliador.readers import read_c6, read_ofx, read_pdf, read_sicoob, read_table


def workbook(rows):
    book = Workbook()
    sheet = book.active
    sheet.title = "Lancamentos"
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(list(row) + [None] * (10 - len(row)))
    result = BytesIO()
    book.save(result)
    return result.getvalue()


def rule(code, pattern="PIX RECEBIDO", kind="contains", direction="entrada"):
    return dict(code=code, pattern=pattern, kind=kind, direction=direction)


class DomainTests(unittest.TestCase):
    def test_money_preserves_cents_and_rejects_invalid(self):
        self.assertEqual(money("-R$ 1.234,56"), -123456)
        self.assertEqual(money("0.05"), 5)
        for value in ("NaN", "Infinity", "0.001", "abc"):
            with self.assertRaises(InputError):
                money(value)

    def test_unknown_and_conflicts_go_to_106(self):
        m = Movement(date(2026, 1, 1), "Pix recebido", 100)
        self.assertEqual(classify(m, [])[0], 106)
        self.assertEqual(classify(m, [rule(151), rule(193)])[0], 106)
        self.assertEqual(classify(m, [rule(151), rule(151)])[0], 151)

    def test_exact_rule_overrides_broad_rule_and_direction_is_separate(self):
        m = Movement(date(2026, 1, 1), "Pix recebido", 100)
        self.assertEqual(classify(m, [rule(151), rule(193, kind="exact")])[0], 193)
        self.assertEqual(classify(m, [rule(150, direction="saida")])[0], 106)


class ReaderTests(unittest.TestCase):
    def test_sicoob_multiline_amount_year_rollover_and_balance(self):
        text = """PERÍODO: 01/01/2026 - 31/01/2026
CONTA: 123.456-7 / EMPRESA EXEMPLO
31/12 SALDO ANTERIOR 100,00
C
02/01 CR CMP VISA 1.234,56
C
Descrição em outra linha
DOC.: 123
02/01 SALDO DO DIA 1.334,56
C
RESUMO
"""
        result = read_sicoob(text)
        self.assertEqual(len(result.movements), 1)
        self.assertEqual(result.movements[0].date, date(2026, 1, 2))
        self.assertEqual(result.movements[0].cents, 123456)
        self.assertEqual(result.movements[0].description, "CR CMP VISA DESCRIÇÃO EM OUTRA LINHA - 123")
        with self.assertRaises(InputError):
            read_sicoob(text.replace("1.334,56", "1.334,57"))

    def test_sicoob_bad_line_is_not_silently_dropped(self):
        text = "PERÍODO: 01/01/2026 - 31/01/2026\n02/01 PIX VALOR DESCONHECIDO\nDOC.: 1\nRESUMO"
        with self.assertRaises(InputError):
            read_sicoob(text)

    def test_c6_uses_launch_date_and_checks_monthly_totals(self):
        text = """Período • 1 de agosto de 2026 até 31 de agosto de 2026
Agência: 1 • Conta: 12345
Agosto 2026 ( 01/08/2026 - 31/08/2026 ) Entradas: R$ 0,00 • Saídas: R$ 0,05
04/08 03/08 Outros gastos IOF CHEQUE ESPECIAL -R$ 0,05
"""
        result = read_c6(text)
        self.assertEqual(result.movements[0].date, date(2026, 8, 4))
        self.assertEqual(result.movements[0].book_date, date(2026, 8, 3))
        with self.assertRaises(InputError):
            read_c6(text.replace("Saídas: R$ 0,05", "Saídas: R$ 0,06"))
        with self.assertRaises(InputError):
            read_c6(text + "05/08 05/08 Tipo desconhecido EXEMPLO R$ 10,00\n")

    def test_encrypted_pdf_wrong_password(self):
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.encrypt("senha-ficticia-de-teste")
        content = BytesIO()
        writer.write(content)
        with self.assertRaisesRegex(InputError, "senha correta"):
            read_pdf(content.getvalue(), "incorreta")

    def test_ofx_xml_and_sgml(self):
        for close in ("</DTPOSTED>", ""):
            content = f"<OFX><BANKACCTFROM><ACCTID>123</ACCTID></BANKACCTFROM><BANKTRANLIST><STMTTRN><DTPOSTED>20260804120000{close}\n<TRNAMT>-12.34\n<FITID>abc\n<MEMO>Pagamento exemplo\n</STMTTRN></BANKTRANLIST></OFX>".encode()
            result = read_ofx(content)
            self.assertEqual(result.movements[0].cents, -1234)
            self.assertEqual(result.bank_account, "123")

    def test_csv_column_mapping_and_signed_values(self):
        result = read_table("Dia;Texto;Quantia\n04/08/2026;Pagamento;-1.234,56\n".encode(), ".csv", {"date": "Dia", "description": "Texto", "value": "Quantia"})
        self.assertEqual(result.movements[0].cents, -123456)
        with self.assertRaises(InputError):
            read_table(b"Data;Descricao;Valor\n04/08/2026;Teste;1.2345\n", ".csv", {})

    def test_xlsx_formulas_are_rejected(self):
        content = workbook([["04/08/2026", "=1+1", 10, 12, 151, 52]])
        with self.assertRaisesRegex(InputError, "fórmulas"):
            read_table(content, ".xlsx", {})


    def test_reconciled_workbook_cannot_be_mistaken_for_statement(self):
        content = workbook([["04/08/2026", "Exemplo fictício", 10, 12, 151, 52]])
        with self.assertRaisesRegex(InputError, "Gerenciar regras"):
            read_table(content, ".xlsx", {})

    def test_chart_pdf_is_directed_to_settings(self):
        with patch("conciliador.readers.PdfReader") as reader:
            reader.return_value.is_encrypted = False
            reader.return_value.pages = [unittest.mock.Mock()]
            reader.return_value.pages[0].extract_text.return_value = "Plano de Contas\nConta S/A Código Histórico D/C Tipo de Saldo"
            with self.assertRaisesRegex(InputError, "Clientes e configurações"):
                read_pdf(b"pdf-ficticio")


class MakroTests(unittest.TestCase):
    def test_private_chart_validation(self):
        content = (ROOT / "assets" / "plano_de_contas.csv").read_bytes()
        prepared, count = prepare_private_chart(content, "plano.csv")
        self.assertGreater(count, 0)
        self.assertIn(b"106", prepared)
        with self.assertRaises(InputError):
            prepare_private_chart(content + content.decode("utf-8-sig").splitlines()[-1].encode() + b"\n", "plano.csv")
        with self.assertRaises(InputError):
            prepare_private_chart(b"conta;codigo\n1;106\n", "plano.csv")

    def test_export_has_exact_schema_and_safe_text(self):
        rows = [dict(date="2026-08-04", description="=EXEMPLO()", cents=-12345, code=106)]
        book = load_workbook(export_workbook(rows, 12), data_only=False)
        self.assertEqual(book.sheetnames, ["Lancamentos"])
        sheet = book.active
        self.assertEqual(list(next(sheet.values)), HEADERS)
        self.assertEqual(sheet.max_row, 2)
        values = list(sheet.values)[1]
        self.assertEqual(values[2:], (123.45, 106, 12, 52, None, None, None, None))
        self.assertEqual(sheet["B2"].data_type, "s")
        self.assertEqual(sheet["A2"].number_format, "[$-416]d/m/yyyy")
        self.assertEqual(sheet["C2"].number_format, "0.00")
        self.assertEqual(sheet["I2"].number_format, "@")

    def test_learning_preserves_conflicts_and_pending_106(self):
        content = workbook([["04/08/2026", "Pagamento", 10, 106, 12, 52], ["04/08/2026", "Pagamento", 10, 150, 12, 52]])
        rules, count = learn_workbook(content, 12)
        self.assertEqual(count, 2)
        self.assertEqual({r[3] for r in rules}, {106, 150})
        with self.assertRaises(InputError):
            learn_workbook(content, 98)

    def test_fractional_inactive_and_synthetic_codes_are_rejected(self):
        for value in (12.5, "12.5", 7, 10000):
            with self.assertRaises(InputError):
                validate_code(value)
        self.assertEqual(validate_code(12.0), 12)


class WebTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.app = create_app({"TESTING": True, "SECRET_KEY": "test-key", "DATABASE": str(Path(self.tmp.name) / "app.sqlite3")})
        self.client = self.app.test_client()
        self.client.get("/")
        with self.client.session_transaction() as session:
            self.csrf = session["csrf"]

    def tearDown(self):
        self.tmp.cleanup()

    def post(self, path, data):
        data["csrf"] = self.csrf
        return self.client.post(path, data=data, follow_redirects=True)

    def create_account(self):
        self.post("/clients", {"name": "Empresa fictícia"})
        self.post("/accounts", {"client_id": "1", "name": "Sicoob", "number": "123-4", "bank_code": "12", "default_in": "106", "default_out": "106"})

    def test_csrf_and_empty_home(self):
        self.assertEqual(self.client.post("/clients", data={"name": "X"}).status_code, 400)
        self.assertIn("Comece por um cliente".encode(), self.client.get("/").data)

    def test_chart_upload_is_saved_only_in_private_path(self):
        path = Path(self.tmp.name) / "data" / "plano_de_contas.csv"
        content = (ROOT / "assets" / "plano_de_contas.csv").read_bytes()
        with patch("app.PRIVATE_CHART", path):
            response = self.post("/chart", {"file": (BytesIO(content), "plano.csv")})
        self.assertIn("Plano de contas salvo".encode(), response.data)
        self.assertTrue(path.exists())

    def test_end_to_end_review_export_period_and_duplicate(self):
        self.create_account()
        response = self.post("/accounts/1/learn", {"file": (BytesIO(workbook([["04/08/2026", "Recebimento", 100, 12, 151, 52]])), "modelo.xlsx")})
        self.assertIn("regras exatas".encode(), response.data)
        content = b"Data;Descricao;Valor\n04/08/2026;Recebimento;100\n05/08/2026;Pagamento;-20\n01/09/2026;Outro;999\n"
        response = self.post("/import", {"account_id": "1", "file": (BytesIO(content), "extrato.csv"), "start": "2026-08-01", "end": "2026-08-31"})
        self.assertEqual(response.status_code, 200)
        self.assertIn("Revise e baixe o Excel".encode(), response.data)
        self.assertNotIn(b"999,00", response.data)
        data = {"code_1": "151", "include_1": "on", "code_2": "106", "include_2": "on"}
        export = self.post("/batches/1/export", data)
        self.assertEqual(export.status_code, 200)
        values = list(load_workbook(BytesIO(export.data), data_only=True).active.values)
        self.assertEqual(values[1][3:6], (12, 151, 52))
        self.assertEqual(values[2][3:6], (106, 12, 52))
        response = self.post("/import", {"account_id": "1", "file": (BytesIO(content), "extrato.csv"), "start": "2026-08-01", "end": "2026-08-31"})
        self.assertIn("já foram importados".encode(), response.data)
        overlap = content + b"02/09/2026;Adicional;1\n"
        response = self.post("/import", {"account_id": "1", "file": (BytesIO(overlap), "outro.csv"), "start": "2026-08-01", "end": "2026-08-31"})
        self.assertIn("Possível repetido".encode(), response.data)
        self.assertIn(b'name="include_3" aria-label="Exportar movimento 1" >', response.data)

    def test_rules_are_scoped_to_account(self):
        self.create_account()
        self.post("/clients", {"name": "Outro cliente"})
        self.post("/accounts", {"client_id": "2", "name": "C6", "number": "5678", "bank_code": "98"})
        self.post("/accounts/1/rules", {"kind": "contains", "direction": "entrada", "pattern": "recebimento", "code": "151"})
        response = self.post("/import", {"account_id": "2", "file": (BytesIO(b"Data;Descricao;Valor\n04/08/2026;Recebimento;100\n"), "extrato.csv")})
        self.assertIn(b'value="106" aria-label="Contrapartida do movimento 1"', response.data)

    def test_bank_account_mismatch_is_rejected(self):
        self.create_account()
        content = b"<OFX><ACCTID>9999\n<STMTTRN><DTPOSTED>20260804\n<TRNAMT>10\n<MEMO>Exemplo\n</STMTTRN></OFX>"
        response = self.post("/import", {"account_id": "1", "file": (BytesIO(content), "extrato.ofx")})
        self.assertIn("difere da conta selecionada".encode(), response.data)

    def test_pdf_password_is_not_retained_in_saved_filename(self):
        self.create_account()
        statement = Statement("PDF de teste", [Movement(date(2026, 8, 4), "Exemplo", 100)], "1234")
        with patch("app.read_statement", return_value=statement):
            response = self.post("/import", {"account_id": "1", "file": (BytesIO(b"pdf-ficticio"), "extrato-senha-ficticia.pdf"), "password": "senha-ficticia"})
        self.assertIn(b"extrato-[protegido].pdf", response.data)
        self.assertNotIn(b"senha-ficticia", response.data)

    def test_first_use_goes_from_setup_to_statement_not_rules(self):
        response = self.client.get("/new", follow_redirects=True)
        self.assertTrue(response.request.path.endswith("/setup"))
        response = self.post("/setup", {"client_name": "Cliente fictício A", "account_name": "Banco fictício", "number": "123-4", "bank_code": "12"})
        self.assertEqual(response.request.path, "/new")
        self.assertIn("Envie o extrato bancário".encode(), response.data)
        with sqlite3.connect(self.app.config["DATABASE"]) as db:
            self.assertEqual(db.execute("SELECT COUNT(*) FROM clients").fetchone()[0], 1)
            self.assertEqual(db.execute("SELECT COUNT(*) FROM batches").fetchone()[0], 0)

    def test_customer_and_account_are_selected_before_upload(self):
        self.create_account()
        response = self.client.get("/new")
        self.assertIn(b'id="client-picker"', response.data)
        self.assertIn(b'id="account-picker"', response.data)
        self.assertNotIn(b'id="import-form"', response.data)
        response = self.client.get("/new?account_id=1")
        self.assertIn(b'id="import-form"', response.data)
        self.assertIn(b'name="account_id" value="1"', response.data)

    def test_setup_validation_is_atomic(self):
        response = self.post("/setup", {"client_name": "Cliente fictício", "account_name": "Banco", "number": "1234", "bank_code": "106"})
        self.assertIn("Selecione a conta bancária correspondente".encode(), response.data)
        with sqlite3.connect(self.app.config["DATABASE"]) as db:
            self.assertEqual(db.execute("SELECT COUNT(*) FROM clients").fetchone()[0], 0)

    def test_json_import_error_keeps_context_and_returns_edit_link(self):
        self.create_account()
        content = b"<OFX><ACCTID>9999\n<STMTTRN><DTPOSTED>20260804\n<TRNAMT>10\n<MEMO>Exemplo\n</STMTTRN></OFX>"
        response = self.client.post("/import", data={"csrf": self.csrf, "account_id": "1", "file": (BytesIO(content), "extrato.ofx")}, headers={"Accept": "application/json"})
        self.assertEqual(response.status_code, 422)
        self.assertEqual(response.json["account_url"], "/accounts/1/edit")
        self.assertIn("9999", response.json["error"])
        with sqlite3.connect(self.app.config["DATABASE"]) as db:
            self.assertEqual(db.execute("SELECT COUNT(*) FROM batches").fetchone()[0], 0)

    def test_edit_wrong_account_number_without_recreating_client(self):
        self.create_account()
        response = self.post("/accounts/1/edit", {"name": "Banco", "number": "567-8", "bank_code": "12"})
        self.assertEqual(response.request.path, "/new")
        self.assertIn(b"567-8", response.data)
        with sqlite3.connect(self.app.config["DATABASE"]) as db:
            self.assertEqual(db.execute("SELECT COUNT(*) FROM accounts").fetchone()[0], 1)

    def test_cannot_reassign_an_account_with_existing_movements(self):
        self.create_account()
        self.post("/import", {"account_id": "1", "file": (BytesIO(b"Data;Descricao;Valor\n04/08/2026;Exemplo;100\n"), "extrato.csv")})
        self.post("/accounts/1/edit", {"name": "Banco", "number": "567-8", "bank_code": "12"})
        with sqlite3.connect(self.app.config["DATABASE"]) as db:
            self.assertEqual(db.execute("SELECT number FROM accounts WHERE id=1").fetchone()[0], "123-4")
            self.assertEqual(db.execute("SELECT COUNT(*) FROM movements").fetchone()[0], 1)

    def test_export_status_and_subsequent_changes(self):
        self.create_account()
        self.post("/import", {"account_id": "1", "file": (BytesIO(b"Data;Descricao;Valor\n04/08/2026;Exemplo;100\n"), "extrato.csv")})
        response = self.client.post("/batches/1/export", data={"csrf": self.csrf, "code_1": "106", "include_1": "on"}, headers={"Accept": "application/json"})
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response.content_type)
        self.assertIn("Excel já gerado".encode(), self.client.get("/").data)
        self.post("/batches/1/review", {"code_1": "151", "include_1": "on"})
        with sqlite3.connect(self.app.config["DATABASE"]) as db:
            self.assertIsNone(db.execute("SELECT exported_at FROM batches WHERE id=1").fetchone()[0])

    def test_migration_preserves_existing_batches_and_movements(self):
        self.create_account()
        self.post("/import", {"account_id": "1", "file": (BytesIO(b"Data;Descricao;Valor\n04/08/2026;Exemplo;100\n"), "extrato.csv")})
        with sqlite3.connect(self.app.config["DATABASE"]) as db:
            db.execute("ALTER TABLE batches DROP COLUMN exported_at")
        self.assertEqual(self.client.get("/").status_code, 200)
        with sqlite3.connect(self.app.config["DATABASE"]) as db:
            self.assertEqual(db.execute("SELECT COUNT(*) FROM movements").fetchone()[0], 1)
            self.assertIn("exported_at", [r[1] for r in db.execute("PRAGMA table_info(batches)")])


if __name__ == "__main__":
    unittest.main()
