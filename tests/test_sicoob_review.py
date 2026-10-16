"""Regressões de leitura, aprendizado e revisão de lotes grandes; dados fictícios."""

from datetime import date
from io import BytesIO
from pathlib import Path
import sqlite3
import tempfile
import unittest
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook

from app import create_app
from conciliador.domain import InputError, Movement, Statement, RuleClassifier, classify
from conciliador.makro import HEADERS, chart, export_workbook, learn_workbook
from conciliador.readers import read_pdf
from conciliador.sicoob_portal import read_sicoob_portal


def portal_sample():
    return """SICOOB — SISBR
EXTRATO DE CONTA CORRENTE
Cooperativa: 0001-0 / SICOOB DEMONSTRACAO
Conta: 000.001-1 / EMPRESA FICTICIA
Periodo: 01/04/2026 - 30/04/2026
HISTÓRICO DE MOVIMENTAÇÃO
Data DocumentoHistórico Valor
02/04102
 DÉB.TRANSF.CONTAS DIF.TITULARIDADE R$ 50,00D
FAV.: PESSOA FICTICIA
01/05/2026, 10:30 Sicoob | Internet Banking
https://ib.sicoob.com.br/sicoobnet/ib/#/home-extrato 1/2
Data DocumentoHistórico Valor
COMPLEMENTO EXEMPLO
02/04REAP
 LIBERAÇÃO DE DEPÓSITO BLOQUEADO R$ 10,00C
02/0488
 DEVOL.CHEQUE DEPOSITADO R$ 10,00D
02/04
 SALDO DO DIA R$ 250,00C
01/04Pix
 PIX RECEBIDO OUTRA IF
Recebimento Pix 00.000.000 0001-00 EMPRESA EXEMPLO R$ 200,00C
01/04
 SALDO DO DIA R$ 300,00C
31/03 SALDO BLOQUEADO ANTERIOR R$ 0,00*
31/03
 SALDO ANTERIOR R$ 100,00C
RESUMO
Saldo em conta: 250,00C
Saldo disponível: 250,00C
01/05/2026, 10:30 Sicoob | Internet Banking
https://ib.sicoob.com.br/sicoobnet/ib/#/home-extrato 2/2
"""


def learning_file(description="TED INTERNET", *, formatted_extras=False, blank_first=False):
    book = Workbook()
    sheet = book.active
    sheet.title = "Lancamentos"
    if blank_first:
        sheet.append([None] * 10)
    sheet.append(HEADERS)
    sheet.append([date(2026, 4, 1), description, 10, 150, 12, 52])
    if formatted_extras:
        sheet.cell(sheet.max_row, 50).number_format = "0.00"
    content = BytesIO()
    book.save(content)
    return content.getvalue()


class SicoobPortalTests(unittest.TestCase):
    def test_reverse_days_join_page_continuation_and_check_all_balances(self):
        statement = read_sicoob_portal(portal_sample())
        self.assertEqual(statement.bank_account, "000.001-1")
        self.assertEqual([m.cents for m in statement.movements], [20000, -5000, 1000, -1000])
        self.assertEqual(statement.movements[0].date, date(2026, 4, 1))
        self.assertEqual(statement.movements[1].document, "102")
        self.assertIn("FAV.: PESSOA FICTICIA COMPLEMENTO EXEMPLO", statement.movements[1].description)
        self.assertNotIn("Internet Banking", statement.movements[1].description)
        self.assertNotIn("https://", statement.movements[1].description)
        self.assertTrue(any("2 saldos diários" in warning for warning in statement.warnings))

    def test_both_known_footer_hosts_and_header_spacing(self):
        text = portal_sample().replace("https://ib.", "https://www.").replace("DocumentoHistórico", "Documento Histórico")
        self.assertEqual(len(read_sicoob_portal(text).movements), 4)

    def test_period_case_accents_and_account_ambiguity(self):
        self.assertEqual(len(read_sicoob_portal(portal_sample().replace("Periodo:", "PERÍODO:")).movements), 4)
        for addition in ["\nConta: 000.002-2 / OUTRA FICTICIA", "\nPeriodo: 01/05/2026 - 31/05/2026"]:
            with self.subTest(addition=addition), self.assertRaisesRegex(InputError, "conta ou período"):
                read_sicoob_portal(portal_sample() + addition)

    def test_unknown_money_line_reports_page_day_and_recovery_action(self):
        text = portal_sample().replace("DEVOL.CHEQUE DEPOSITADO R$ 10,00D", "DEVOL.CHEQUE DEPOSITADO R$ ILEGIVEL")
        with self.assertRaises(InputError) as error:
            read_sicoob_portal(text)
        message = str(error.exception)
        for detail in ["Sicoob Internet Banking", "página 2", "02/04", "C/D", "OFX"]:
            self.assertIn(detail, message)

    def test_amount_without_sign_or_second_amount_cannot_be_skipped(self):
        for altered in ["R$ 10,00", "R$ 10,00D R$ 0,00C"]:
            with self.subTest(altered=altered), self.assertRaises(InputError):
                read_sicoob_portal(portal_sample().replace("DEVOL.CHEQUE DEPOSITADO R$ 10,00D", "DEVOL.CHEQUE DEPOSITADO " + altered))

    def test_missing_page_and_duplicate_page_are_reported(self):
        for text in [portal_sample().replace("2/2", "3/3").replace("1/2", "1/3"), portal_sample().replace("2/2", "1/2")]:
            with self.subTest(text=text), self.assertRaisesRegex(InputError, "página"):
                read_sicoob_portal(text)

    def test_daily_final_and_missing_opening_balance_all_abort(self):
        variants = [portal_sample().replace("SALDO DO DIA R$ 250,00C", "SALDO DO DIA R$ 250,01C"), portal_sample().replace("Saldo em conta: 250,00C", "Saldo em conta: 250,01C"), portal_sample().replace("31/03\n SALDO ANTERIOR R$ 100,00C", "")]
        for text in variants:
            with self.subTest(text=text), self.assertRaises(InputError):
                read_sicoob_portal(text)

    def test_dates_outside_declared_period_are_rejected(self):
        with self.assertRaisesRegex(InputError, "data"):
            read_sicoob_portal(portal_sample().replace("02/04", "02/05"))

    def test_legitimate_equal_transactions_are_retained(self):
        text = portal_sample().replace("02/0488\n DEVOL.CHEQUE DEPOSITADO R$ 10,00D", "02/0488\n DEVOL.CHEQUE DEPOSITADO R$ 5,00D\n02/0488\n DEVOL.CHEQUE DEPOSITADO R$ 5,00D")
        result = read_sicoob_portal(text)
        self.assertEqual([m.cents for m in result.movements], [20000, -5000, 1000, -500, -500])
        self.assertEqual(result.movements[-1], result.movements[-2])

    def test_printed_blocked_deposit_is_excluded_but_its_release_is_retained(self):
        text = portal_sample().replace("02/04REAP", "02/04REAP DEP.CHEQUE BLOQ.1D R$ 10,00*\n02/04REAP")
        statement = read_sicoob_portal(text)
        self.assertEqual(len(statement.movements), 4)
        self.assertTrue(any("2 registros" in w for w in statement.warnings))

    def test_dispatch_and_old_layout_can_coexist(self):
        with patch("conciliador.readers.PdfReader") as pdf:
            pdf.return_value.is_encrypted = False
            pdf.return_value.pages = [Mock(extract_text=Mock(return_value=portal_sample()))]
            self.assertIn("Internet Banking", read_pdf(b"ficticio").layout)


class LearningAndClassificationTests(unittest.TestCase):
    def test_leading_empty_row_and_formatted_empty_extra_columns_are_accepted(self):
        rules, count = learn_workbook(learning_file(formatted_extras=True, blank_first=True), 12)
        self.assertEqual(count, 1)
        self.assertEqual(rules, [("exact", "TED INTERNET", "saida", 150)])

    def test_nonempty_extra_column_is_not_silently_ignored(self):
        book = load_workbook(BytesIO(learning_file(formatted_extras=True)))
        book.active.cell(2, 50, "DADO FORA DO MODELO")
        content = BytesIO(); book.save(content)
        with self.assertRaisesRegex(InputError, "Linha 2.*dez colunas"):
            learn_workbook(content.getvalue(), 12)

    def test_learning_and_export_load_chart_once_for_the_whole_workbook(self):
        with patch("conciliador.makro.chart", wraps=chart) as read_chart:
            rules, count = learn_workbook(learning_file(), 12)
            self.assertEqual(read_chart.call_count, 1)
        rows = [dict(date="2026-04-01", description="EXEMPLO", cents=-100, code=106)] * 500
        with patch("conciliador.makro.chart", wraps=chart) as read_chart:
            output = export_workbook(rows, 12)
            self.assertEqual(load_workbook(output).active.max_row, 501)
            self.assertEqual(read_chart.call_count, 1)

    def test_known_presentation_variants_match_exact_text_and_document_only(self):
        movement = read_sicoob_portal(portal_sample()).movements[1]
        patterns = [movement.description, *movement.description_variants]
        for pattern in patterns:
            with self.subTest(pattern=pattern):
                rule = dict(kind="exact", direction="saida", pattern=pattern, code=150)
                self.assertEqual(classify(movement, [rule])[0], 150)
                changed = dict(rule, pattern=pattern.replace("PESSOA FICTICIA", "OUTRA PESSOA"))
                self.assertEqual(classify(movement, [changed])[0], 106)

    def test_variant_conflicts_go_to_106_and_never_override_exact_with_contains(self):
        movement = read_sicoob_portal(portal_sample()).movements[1]
        exact = dict(kind="exact", direction="saida", pattern=movement.description_variants[0], code=150)
        broad = dict(kind="contains", direction="saida", pattern="DÉB.TRANSF.", code=251)
        self.assertEqual(classify(movement, [exact, broad])[0], 150)
        conflict = dict(exact, pattern=movement.description_variants[1], code=251)
        self.assertEqual(classify(movement, [exact, conflict])[0], 106)
        self.assertEqual(classify(movement, [dict(exact, direction="entrada")])[0], 106)

    def test_truncated_reference_is_not_guessed(self):
        movement = read_sicoob_portal(portal_sample()).movements[1]
        rule = dict(kind="exact", direction="saida", pattern="DÉB.TRANSF.CONTAS DIF.TITULARIDADE", code=150)
        self.assertEqual(RuleClassifier([rule]).classify(movement)[0], 106)


class LargeBatchAndErrorTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.app = create_app(dict(TESTING=True, SECRET_KEY="fictional-key", DATABASE=str(Path(self.tmp.name) / "app.sqlite3")))
        self.client = self.app.test_client();self.client.get("/")
        with self.client.session_transaction() as session:
            self.csrf = session["csrf"]
        self.client.post("/clients", data=dict(csrf=self.csrf, name="Empresa fictícia"))
        self.client.post("/accounts", data=dict(csrf=self.csrf, client_id=1, name="Sicoob", number="000001-1", bank_code=12))

    def tearDown(self):
        self.tmp.cleanup()

    def test_full_2600_row_review_export_and_repeat_keep_every_transaction(self):
        csv = "Data;Descricao;Valor\n" + "01/04/2026;FICTICIO;1,01\n" * 2600
        response = self.client.post("/import", data=dict(csrf=self.csrf, account_id=1, file=(BytesIO(csv.encode()), "extrato.csv")), headers={"Accept": "application/json"})
        self.assertEqual(response.status_code, 200)
        data = dict(csrf=self.csrf)
        for movement_id in range(1, 2601):
            data[f"code_{movement_id}"] = "151"
            data[f"include_{movement_id}"] = "on"
        response = self.client.post("/batches/1/export", data=data, content_type="multipart/form-data", headers={"Accept": "application/json"})
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.data), data_only=True).active
        self.assertEqual(sheet.max_row, 2601)
        self.assertEqual(sheet.max_column, 10)
        self.assertEqual(list(next(sheet.values)), HEADERS)
        self.assertTrue(all(row[3:6] == (12, 151, 52) and row[6:] == (None,) * 4 for row in list(sheet.values)[1:]))
        repeated = self.client.post("/import", data=dict(csrf=self.csrf, account_id=1, file=(BytesIO(csv.encode()), "extrato.csv")), headers={"Accept": "application/json"})
        self.assertEqual(repeated.json["redirect"], "/batches/1")
        with sqlite3.connect(self.app.config["DATABASE"]) as db:
            self.assertEqual(db.execute("SELECT COUNT(*),SUM(include_row),SUM(cents) FROM movements").fetchone(), (2600, 2600, 262600))
            self.assertEqual(db.execute("SELECT COUNT(*) FROM batches").fetchone()[0], 1)

    def test_form_limit_error_identifies_review_instead_of_claiming_file_size(self):
        self.app.config["MAX_FORM_PARTS"] = 2
        response = self.client.post("/batches/1/export", data=dict(csrf=self.csrf, code_1="106", include_1="on"), content_type="multipart/form-data", headers={"Accept": "application/json"})
        self.assertEqual(response.status_code, 413)
        self.assertIn("revisão", response.json["error"])
        self.assertNotIn("Arquivo maior que 16 MB", response.json["error"])

    def test_pdf_reader_error_does_not_offer_unrelated_account_edit(self):
        with patch("app.read_statement", side_effect=InputError("Sicoob — página 2: sinal C/D ausente. Envie OFX.")):
            response = self.client.post("/import", data=dict(csrf=self.csrf, account_id=1, file=(BytesIO(b"ficticio"), "extrato.pdf")), headers={"Accept": "application/json"})
        self.assertEqual(response.status_code, 422)
        self.assertIn("página 2", response.json["error"])
        self.assertNotIn("account_url", response.json)

    def test_oversized_movement_selection_has_a_period_action_without_saving_partial_batch(self):
        movement = Movement(date(2026, 4, 1), "FICTICIO", 100)
        with patch("app.read_statement", return_value=Statement("Ficticio", [movement] * 20001)):
            response = self.client.post("/import", data=dict(csrf=self.csrf, account_id=1, file=(BytesIO(b"ficticio"), "extrato.pdf")), headers={"Accept": "application/json"})
        self.assertEqual(response.status_code, 422)
        self.assertIn("De/Até", response.json["error"])
        self.assertIn("20001", response.json["error"])
        with sqlite3.connect(self.app.config["DATABASE"]) as db:
            self.assertEqual(db.execute("SELECT COUNT(*) FROM batches").fetchone()[0], 0)

    def test_pdf_page_extraction_error_identifies_page_and_alternative(self):
        with patch("conciliador.readers.PdfReader") as pdf:
            pdf.return_value.is_encrypted = False
            pdf.return_value.pages = [Mock(extract_text=Mock(return_value="first page")), Mock(extract_text=Mock(side_effect=ValueError("bad-page")))]
            with self.assertRaisesRegex(InputError, "página 2.*OFX"):
                read_pdf(b"ficticio")


if __name__ == "__main__":
    unittest.main()
