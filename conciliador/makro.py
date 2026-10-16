from copy import copy
from datetime import datetime
from io import BytesIO, StringIO
from pathlib import Path
import csv
import re
from decimal import Decimal

from openpyxl import load_workbook
from pypdf import PdfReader

from .domain import InputError, normalize

ROOT = Path(__file__).resolve().parent.parent
HEADERS = ["Data", "Conta", "Valor", "Debito", "Credito", "Historico", "SCP", "Departamento", "Documento", "CNPJ"]
CHART_FIELDS = ["conta", "tipo", "codigo", "descricao", "natureza", "tipo_saldo", "inativada"]
PRIVATE_CHART = ROOT / "data" / "plano_de_contas.csv"


def chart():
    path = PRIVATE_CHART if PRIVATE_CHART.exists() else ROOT / "assets" / "plano_de_contas.csv"
    with path.open(encoding="utf-8-sig") as file:
        return {int(row["codigo"]): row for row in csv.DictReader(file, delimiter=";")}


def prepare_private_chart(content, filename):
    """Validate a chart for local storage; never add uploaded charts to assets."""
    if filename.lower().endswith(".pdf"):
        reader = PdfReader(BytesIO(content))
        if reader.is_encrypted or len(reader.pages) > 200:
            raise InputError("Envie o plano sem senha, com até 200 páginas.")
        text = "\n".join(page.extract_text(extraction_mode="layout") or "" for page in reader.pages)
        candidates = [line for line in text.splitlines() if re.match(r"^\s*\d+\.\d+\.\d+\.\d+\.\d+\s", line)]
        pattern = re.compile(r"^\s*(\d+\.\d+\.\d+\.\d+\.\d+)\s+([SA])\s+(\d+)\s+(.+?)\s+([DC])\s+(DC|D|C)\s*$")
        rows = []
        for line in candidates:
            match = pattern.fullmatch(line)
            if not match:
                raise InputError("O plano contém uma linha fora do layout reconhecido; nenhuma conta foi salva.")
            values = list(match.groups()) + ["sim" if "(Conta Inativada)" in match[4] else "nao"]
            rows.append(dict(zip(CHART_FIELDS, values)))
    elif filename.lower().endswith(".csv"):
        reader = csv.DictReader(StringIO(content.decode("utf-8-sig")), delimiter=";")
        if reader.fieldnames != CHART_FIELDS:
            raise InputError("CSV deve ter as colunas: " + ";".join(CHART_FIELDS))
        rows = list(reader)
    else:
        raise InputError("Envie o plano de contas em PDF do Makro ou CSV com o cabeçalho documentado.")
    if not rows or len(rows) > 20000:
        raise InputError("O plano deve conter entre 1 e 20.000 contas.")
    seen_codes, seen_accounts = set(), set()
    for row in rows:
        if not re.fullmatch(r"\d+", row.get("codigo") or "") or not re.fullmatch(r"\d+\.\d+\.\d+\.\d+\.\d+", row.get("conta") or ""):
            raise InputError("Código ou classificação contábil inválidos no plano.")
        if row["tipo"] not in ("S", "A") or row["inativada"] not in ("sim", "nao") or not row["descricao"]:
            raise InputError("Tipo, descrição ou situação inválidos no plano.")
        code = int(row["codigo"])
        if code <= 0 or code in seen_codes or row["conta"] in seen_accounts:
            raise InputError("O plano contém códigos/classificações duplicados ou inválidos.")
        seen_codes.add(code)
        seen_accounts.add(row["conta"])
    if not any(int(r["codigo"]) == 106 and r["tipo"] == "A" and r["inativada"] == "nao" for r in rows):
        raise InputError("O plano precisa incluir a conta analítica ativa 106 para pendências.")
    output = StringIO()
    writer = csv.DictWriter(output, fieldnames=CHART_FIELDS, delimiter=";")
    writer.writeheader()
    writer.writerows(rows)
    return output.getvalue().encode("utf-8-sig"), len(rows)


def validate_code(value, bank_code=None, *, accounts=None):
    try:
        raw = str(value).strip()
        if not re.fullmatch(r"\d+(?:\.0+)?", raw):
            raise ValueError()
        code = int(Decimal(raw))
    except (TypeError, ValueError) as exc:
        raise InputError("Informe um código contábil numérico.") from exc
    row = (accounts if accounts is not None else chart()).get(code)
    if not row or row["tipo"] != "A" or row["inativada"] == "sim":
        raise InputError(f"Código {code} não é uma conta analítica ativa no plano fornecido.")
    if bank_code is not None and code == int(bank_code):
        raise InputError("A contrapartida deve ser diferente da conta bancária.")
    return code


def learn_workbook(content, bank_code):
    book = load_workbook(BytesIO(content), read_only=True, data_only=False)
    try:
        if book.sheetnames != ["Lancamentos"]:
            raise InputError("A planilha de aprendizado deve ter apenas a aba Lancamentos.")
        rows = book.active.iter_rows(values_only=True)
        header_found = False
        rules = set()
        count = 0
        accounts = chart()
        for n, row in enumerate(rows, 1):
            if all(x is None for x in row):
                continue
            if not header_found:
                if list(row[:10]) != HEADERS or any(value is not None for value in row[10:]):
                    raise InputError(f"Linha {n}: cabeçalho diferente do modelo Makro. Use Data, Conta, Valor, Debito, Credito, Historico, SCP, Departamento, Documento e CNPJ, nessa ordem. Linhas iniciais vazias e colunas extras sem dados são aceitas no aprendizado.")
                header_found = True
                continue
            if any(value is not None for value in row[10:]):
                raise InputError(f"Linha {n}: há dados além das dez colunas Makro. Remova esses dados ou copie somente as dez colunas para uma nova planilha.")
            if count >= 20000:
                raise InputError("Planilha excede 20.000 linhas.")
            if row[5] != 52 or not row[1]:
                raise InputError(f"Linha {n}: histórico deve ser 52 e a descrição deve estar preenchida.")
            debit, credit = validate_code(row[3], accounts=accounts), validate_code(row[4], accounts=accounts)
            if (debit == bank_code) == (credit == bank_code):
                raise InputError(f"Linha {n}: a conta bancária {bank_code} deve aparecer em exatamente um lado.")
            direction = "entrada" if debit == bank_code else "saida"
            counterpart = credit if direction == "entrada" else debit
            rules.add(("exact", normalize(row[1]), direction, counterpart))
            count += 1
        if not count:
            raise InputError("Planilha de aprendizado sem lançamentos.")
        return sorted(rules), count
    finally:
        book.close()


def export_workbook(rows, bank_code):
    if not rows:
        raise InputError("Selecione ao menos um movimento para exportar.")
    book = load_workbook(ROOT / "assets" / "makro.xlsx")
    sheet = book["Lancamentos"]
    styles = [copy(sheet.cell(2, c)._style) for c in range(1, 11)]
    sheet.delete_rows(2, sheet.max_row)
    accounts = chart()
    for n, row in enumerate(rows, 2):
        code = validate_code(row["code"], bank_code, accounts=accounts)
        amount = abs(row["cents"]) / 100
        incoming = row["cents"] >= 0
        values = [datetime.strptime(row["date"], "%Y-%m-%d"), row["description"], amount,
                  bank_code if incoming else code, code if incoming else bank_code,
                  52, None, None, None, None]
        for col, value in enumerate(values, 1):
            cell = sheet.cell(n, col, value)
            cell._style = copy(styles[col - 1])
            if isinstance(value, str):
                cell.data_type = "s"  # Text from bank files must never become an Excel formula.
    out = BytesIO()
    book.save(out)
    out.seek(0)
    return out
